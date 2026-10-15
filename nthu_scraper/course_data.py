"""Two-stage official course conversion, normalization, and publication."""

import hashlib
import html
import json
import logging
import math
import os
import re
import tempfile
from dataclasses import asdict, dataclass
from dataclasses import field as dataclass_field
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

FIELD_MAPPING = {
    "id": ["科號"],
    "chinese_title": ["課程中文名稱", "中文課名"],
    "english_title": ["課程英文名稱", "英文課名"],
    "credit": ["學分數", "學分"],
    "time": ["上課時間"],
    "classroom": ["教室"],
    "class_room_and_time": ["教室與上課時間"],
    "lecturer": ["授課教師", "教師姓名", "教師"],
    "size_limit": ["人限"],
    "student_count": ["總人數"],
    "language": ["授課語言"],
    "note": ["備註", "備註欄"],
    "suspend": ["停開註記"],
    "limit_note": ["課程限制說明"],
    "freshman_reservation": ["新生保留人數"],
    "object": ["開課對象", "選課限制條件"],
    "ge_type": ["通識對象"],
    "ge_category": ["通識分類", "通識類別"],
    "prerequisite": ["擋修說明"],
    "expertise": ["第一二專長對應"],
    "program": ["學分學程對應"],
    "no_extra_selection": ["不可加簽說明"],
    "required_optional_note": ["必選修說明", "此課程已列入之系所班別"],
}
PUBLIC_TEXT_FIELDS = (
    "id",
    "chinese_title",
    "english_title",
    "note",
    "limit_note",
    "object",
    "ge_type",
    "ge_category",
    "no_extra_selection",
)
CONVERTED_VERSION = 1


@dataclass(frozen=True)
class CourseSource:
    id: str
    url: str
    format: str = "json"
    sheets: tuple[str, ...] = ()

    @property
    def filename(self) -> str:
        return f"{self.id}.{self.format}"

    @property
    def priority(self) -> int:
        return 30 if self.id == "latest" else 20 if self.format == "xlsx" else 10


SOURCES = (
    CourseSource(
        "latest",
        "https://www.ccxp.nthu.edu.tw/ccxp/INQUIRE/JH/OPENDATA/open_course_data.json",
    ),
    CourseSource(
        "10820",
        "https://www.ccxp.nthu.edu.tw/ccxp/INQUIRE/JH/OPENDATA/open_course_data_10820.json",
    ),
    CourseSource(
        "10910-11110",
        "https://curricul.site.nthu.edu.tw/var/file/208/1208/img/474/unicode_1091_1111.json",
    ),
    CourseSource(
        "11120-11220",
        "https://curricul.site.nthu.edu.tw/var/file/208/1208/img/474/11120-11220JSON.json",
    ),
    CourseSource(
        "108-110",
        "https://curricul.site.nthu.edu.tw/var/file/208/1208/img/1657/108-110coursedata_n(3).xlsx",
        "xlsx",
        ("108-110",),
    ),
    CourseSource(
        "111-113",
        "https://curricul.site.nthu.edu.tw/var/file/208/1208/img/1657/111all-113all_courses_redo.xlsx",
        "xlsx",
        ("111全及112全", "11310", "11320"),
    ),
    CourseSource(
        "11410",
        "https://curricul.site.nthu.edu.tw/var/file/208/1208/img/1350/11410_course_data_fin.xlsx",
        "xlsx",
        ("11410",),
    ),
)
SOURCE_BY_ID = {source.id: source for source in SOURCES}
ARCHIVED_LATEST_ID = re.compile(r"latest-(?:\d{5})(?:-\d{5})*$")


def source_for_id(source_id: str) -> CourseSource:
    if source_id in SOURCE_BY_ID:
        return SOURCE_BY_ID[source_id]
    if ARCHIVED_LATEST_ID.fullmatch(source_id):
        return CourseSource(source_id, SOURCE_BY_ID["latest"].url)
    raise ValueError(f"Unknown course archive identity: {source_id}")


EXTRA_FIELDS = {
    "department_code": ("code", "代碼", "開課代碼"),
    "department_name": ("系所全名", "開課單位"),
    "college_name": ("院全名", "開課院別"),
    "course_level_code": ("課程字頭",),
    "course_level": ("課程屬性",),
    "source_semester": ("學年期",),
    "remote_teaching": ("遠距課",),
    "enrollment_after_selection": ("加退選截止後人數(選課資料庫)",),
    "enrollment_after_withdrawal": ("停修後人數",),
    "lecturer_names_zh": ("授課教師中文",),
    "lecturer_names_en": ("授課教師英文",),
}
ADDITIONAL_ALIASES = {
    "note": ("備註欄",),
    "language": ("授課語言(加退選截止)",),
    "expertise": ("專長", "專長相關"),
    "program": ("學分學程",),
    "prerequisite": ("擋修",),
    "required_optional_note": ("課程列入系所班別",),
}


def raw_text(value: Any) -> str:
    return "" if value is None else str(value)


def text(value: Any) -> str:
    return raw_text(value).strip()


def display_text(value: Any) -> str:
    value = re.sub(r"<br\s*/?>", " ", text(value), flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", html.unescape(value)).strip()


def decode_json(body: bytes) -> list[dict[str, Any]]:
    data = json.loads(body)
    if isinstance(data, dict) and "工作表1" in data:
        data = data["工作表1"]
    if not isinstance(data, list) or not data:
        raise ValueError("Course collection must be a non-empty list")
    if not all(isinstance(row, dict) for row in data):
        raise ValueError("Every course record must be an object")
    return data


def course_key(value: Any) -> str:
    identifier = text(value)
    if (
        len(identifier) <= 5
        or not identifier[:5].isascii()
        or not identifier[:5].isdigit()
        or not identifier[5:].strip()
    ):
        raise ValueError(f"Course has no usable identity: {identifier!r}")
    return re.sub(r" +", "", identifier)


@dataclass
class SourceRecord:
    values: dict[str, Any]
    origin: dict[str, Any]
    field_origins: dict[str, dict[str, Any]] = dataclass_field(default_factory=dict)


def read_source(body: bytes, source: CourseSource) -> list[SourceRecord]:
    if source.format == "json":
        rows = decode_json(body)
        if ARCHIVED_LATEST_ID.fullmatch(source.id):
            retained = set(source.id.removeprefix("latest-").split("-"))
            selected = [
                (index, row)
                for index, row in enumerate(rows, 1)
                if course_key(row.get("科號"))[:5] in retained
            ]
            if not selected:
                raise ValueError(
                    f"{source.id}: archived response has no retained courses"
                )
        else:
            selected = list(enumerate(rows, 1))
        return [
            SourceRecord(row, {"source_id": source.id, "row": index})
            for index, row in selected
        ]
    workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
    records = []
    try:
        for name in source.sheets:
            if name not in workbook.sheetnames:
                raise ValueError(f"{source.id}: missing course sheet {name!r}")
            sheet = workbook[name]
            headers = None
            count = 0
            for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                if headers is None:
                    if "科號" in cells:
                        headers = [text(cell) for cell in cells]
                        nonempty = [key for key in headers if key]
                        if len(nonempty) != len(set(nonempty)):
                            raise ValueError(f"{source.id}/{name}: duplicate headers")
                    continue
                if not any(cell is not None for cell in cells):
                    continue
                row = {
                    key: value
                    for key, value in zip(headers, cells, strict=False)
                    if key
                }
                if not text(row.get("科號")):
                    raise ValueError(f"{source.id}/{name}:{number}: missing course ID")
                records.append(
                    SourceRecord(
                        row,
                        {"source_id": source.id, "sheet": name, "row": number},
                    )
                )
                count += 1
            if headers is None or not count:
                raise ValueError(f"{source.id}/{name}: no course records")
        if source.id == "11410":
            sheet_name = "1141110教師data"
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"{source.id}: missing teacher supplement")
            main_records = {course_key(row.values["科號"]): row for row in records}
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            if tuple(next(rows)) != (
                "科號",
                "課程中文名稱",
                "授課教師中文",
                "授課教師英文",
            ):
                raise ValueError("11410: teacher supplement headers changed")
            seen_teachers = {}
            for index, cells in enumerate(rows, 2):
                if not any(cell is not None for cell in cells):
                    continue
                key = course_key(cells[0])
                if key not in main_records:
                    continue
                names = tuple(cells[2:4])
                if key in seen_teachers and seen_teachers[key] != names:
                    raise ValueError(f"11410: conflicting teacher supplement for {key}")
                seen_teachers[key] = names
                row = main_records[key]
                row.values["授課教師中文"], row.values["授課教師英文"] = names
                row.field_origins.update(
                    {
                        field: {
                            "source_id": source.id,
                            "sheet": sheet_name,
                            "row": index,
                        }
                        for field in ("lecturer_names_zh", "lecturer_names_en")
                    }
                )
    finally:
        workbook.close()
    return records


def canonical_fields(
    values: dict[str, Any], mapping: dict[str, list[str]] = FIELD_MAPPING
) -> dict[str, Any]:
    aliases = {
        **{
            key: (*names, *ADDITIONAL_ALIASES.get(key, ()))
            for key, names in mapping.items()
        },
        **EXTRA_FIELDS,
    }
    result = {}
    for field, names in aliases.items():
        for name in names:
            if name in values:
                value = values[name]
                # A present blank capacity means unlimited, not an absent field.
                if field == "size_limit" and value is None:
                    value = ""
                if value is not None:
                    result[field] = value
                break
    return result


def convert_source(body: bytes, source: CourseSource) -> dict[str, Any]:
    """Extract source records without renaming or cleaning their field values."""
    return {
        "schema_version": CONVERTED_VERSION,
        "source_id": source.id,
        "format": source.format,
        "sha256": hashlib.sha256(body).hexdigest(),
        "records": [asdict(record) for record in read_source(body, source)],
    }


def converted_records(document: Any, source: CourseSource) -> list[SourceRecord]:
    """Validate a stage-one document before it enters normalization."""
    if (
        not isinstance(document, dict)
        or set(document)
        != {"schema_version", "source_id", "format", "sha256", "records"}
        or type(document["schema_version"]) is not int
        or document["schema_version"] != CONVERTED_VERSION
        or document["source_id"] != source.id
        or document["format"] != source.format
        or not isinstance(document["sha256"], str)
        or re.fullmatch(r"[a-f0-9]{64}", document["sha256"]) is None
        or not isinstance(document["records"], list)
        or not document["records"]
    ):
        raise ValueError(f"{source.id}: invalid converted course document")
    records = []
    for record in document["records"]:
        if (
            not isinstance(record, dict)
            or set(record) != {"values", "origin", "field_origins"}
            or not isinstance(record["values"], dict)
            or not isinstance(record["field_origins"], dict)
        ):
            raise ValueError(f"{source.id}: invalid converted course record")
        for origin in (record["origin"], *record["field_origins"].values()):
            expected_keys = (
                {"source_id", "row", "sheet"}
                if source.format == "xlsx"
                else {"source_id", "row"}
            )
            if (
                not isinstance(origin, dict)
                or set(origin) != expected_keys
                or origin["source_id"] != source.id
                or type(origin["row"]) is not int
                or origin["row"] < 1
                or (
                    source.format == "xlsx"
                    and origin["sheet"] not in (*source.sheets, "1141110教師data")
                )
            ):
                raise ValueError(f"{source.id}: invalid converted record origin")
        records.append(SourceRecord(**record))
    return records


def merge_sources(
    inputs: dict[str, bytes], mapping: dict[str, list[str]] = FIELD_MAPPING
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    return merge_converted_sources(
        {key: convert_source(body, source_for_id(key)) for key, body in inputs.items()},
        mapping,
    )


def merge_converted_sources(
    inputs: dict[str, dict[str, Any]],
    mapping: dict[str, list[str]] = FIELD_MAPPING,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    records_by_source = {
        key: converted_records(document, source_for_id(key))
        for key, document in inputs.items()
    }
    current_semesters = (
        {
            course_key(record.values.get("科號"))[:5]
            for record in records_by_source["latest"]
        }
        if "latest" in inputs
        else set()
    )
    merged: dict[str, dict[str, Any]] = {}
    manifest: dict[str, Any] = {
        "path_base": "courses",
        "sources": [],
        "conflicts": {},
        "filled_fields": 0,
        "source_repairs": [],
    }
    for source in sorted(
        (source_for_id(key) for key in inputs),
        key=lambda item: (-item.priority, item.id),
    ):
        document = inputs[source.id]
        records = records_by_source[source.id]
        if ARCHIVED_LATEST_ID.fullmatch(source.id):
            records = [
                record
                for record in records
                if course_key(record.values.get("科號"))[:5] not in current_semesters
            ]
        seen: dict[str, dict[str, Any]] = {}
        semesters = set()
        previous_fields: dict[str, Any] = {}
        for record in records:
            fields = canonical_fields(record.values, mapping)
            if (
                source.format == "json"
                and text(fields.get("id")) == '"'
                and text(previous_fields.get("english_title")).startswith('"')
            ):
                # A historical Excel-to-JSON converter split a quoted multiline
                # title into two records. Only accept recovery from intact Excel.
                previous_key = course_key(previous_fields.get("id"))
                authoritative = merged.get(previous_key)
                if (
                    authoritative is None
                    or authoritative["priorities"]["id"] != 20
                    or text(authoritative["values"].get("english_title"))
                    != text(previous_fields["english_title"]).removeprefix('"')
                ):
                    raise ValueError(f"{record.origin}: unrecoverable split JSON row")
                repair = {
                    **record.origin,
                    "course_id": authoritative["values"]["id"],
                    "reason": "Split quoted JSON row recovered from authoritative Excel",
                    "recovered_from": authoritative["field_sources"]["id"],
                }
                manifest["source_repairs"].append(repair)
                logger.warning("Recovered malformed historical JSON row: %s", repair)
                continue
            key = course_key(fields.get("id"))
            if not (
                text(fields.get("chinese_title")) or text(fields.get("english_title"))
            ):
                raise ValueError(f"{record.origin}: no usable course title")
            semesters.add(key[:5])
            if key in seen:
                if seen[key] != fields:
                    raise ValueError(f"{record.origin}: conflicting duplicate {key}")
                continue
            seen[key] = fields
            previous_fields = fields
            current = merged.setdefault(
                key,
                {"values": {}, "field_sources": {}, "sources": [], "priorities": {}},
            )
            supplementing = bool(current["values"])
            current["sources"].append(record.origin)
            for origin in record.field_origins.values():
                if origin not in current["sources"]:
                    current["sources"].append(origin)
            for field, value in fields.items():
                missing = field not in current["values"] or (
                    not text(current["values"][field])
                    and field not in ("size_limit", "suspend", "id")
                    and bool(text(value))
                )
                if missing:
                    current["values"][field] = value
                    current["field_sources"][field] = record.field_origins.get(
                        field, record.origin
                    )
                    current["priorities"][field] = source.priority
                    if supplementing:
                        manifest["filled_fields"] += 1
                elif text(current["values"][field]) != text(value):
                    if current["priorities"][field] == source.priority:
                        raise ValueError(f"{key}: equal-priority conflict in {field}")
                    conflicts = manifest["conflicts"]
                    conflicts[field] = conflicts.get(field, 0) + 1
        manifest["sources"].append(
            {
                "id": source.id,
                "url": source.url,
                "path": f"raw/{source.filename}",
                "sha256": document["sha256"],
                "format": source.format,
                "sheets": list(source.sheets),
                "supplemental_sheets": sorted(
                    {
                        origin["sheet"]
                        for record in records
                        for origin in record.field_origins.values()
                    }
                ),
                "records": len(records),
                "semesters": sorted(semesters),
            }
        )
    return [merged[key] for key in sorted(merged)], manifest


def parse_meetings(values: dict[str, Any], issues: list[str]) -> list[dict[str, Any]]:
    meetings = []
    combined = []
    if "classroom" not in values or "time" not in values:
        for line in raw_text(values.get("class_room_and_time")).splitlines():
            if not line.strip():
                continue
            parts = line.split("\t")
            if len(parts) > 2:
                issues.append("Unexpected extra classroom/time columns")
            combined.append(
                {
                    "classroom": parts[0].strip() or None,
                    "time": parts[1].strip() or None if len(parts) > 1 else None,
                }
            )
    if "classroom" in values or "time" in values:
        rooms = (
            raw_text(values.get("classroom")).splitlines()
            if "classroom" in values
            else []
        )
        times = raw_text(values.get("time")).splitlines() if "time" in values else []
        if not any(room.strip() for room in rooms):
            rooms = []
        if not any(time.strip() for time in times):
            times = []
        if len(rooms) == len(times) == 1 and "," in rooms[0] and "," in times[0]:
            split_rooms = rooms[0].split(",")
            split_times = times[0].split(",")
            if len(split_rooms) == len(split_times):
                rooms, times = split_rooms, split_times
            else:
                issues.append("Unequal comma-separated classroom/time counts")
        elif combined and len(rooms) == 1 and "," in rooms[0] and "time" not in values:
            rooms = rooms[0].split(",")
        elif (
            combined
            and len(times) == 1
            and "," in times[0]
            and "classroom" not in values
        ):
            times = times[0].split(",")
        if rooms and times and len(rooms) != len(times):
            issues.append("Unequal separate classroom/time line counts")
        if combined and rooms and "time" not in values and len(rooms) != len(combined):
            issues.append("Unequal separate classroom/combined meeting counts")
        if (
            combined
            and times
            and "classroom" not in values
            and len(times) != len(combined)
        ):
            issues.append("Unequal separate time/combined meeting counts")
        for index in range(max(len(rooms), len(times), len(combined))):
            meetings.append(
                {
                    "classroom": (rooms[index].strip() or None)
                    if "classroom" in values and index < len(rooms)
                    else combined[index]["classroom"]
                    if "classroom" not in values and index < len(combined)
                    else None,
                    "time": (times[index].strip() or None)
                    if "time" in values and index < len(times)
                    else combined[index]["time"]
                    if "time" not in values and index < len(combined)
                    else None,
                }
            )
    else:
        meetings = combined
    for meeting in meetings:
        raw_time = meeting["time"] or ""
        codes = re.findall(r"[MTWRFSU](?:[0-9]+|[A-Za-z])", raw_time)
        if "".join(codes) != re.sub(r"[\s,]+", "", raw_time):
            issues.append(f"Unrecognized time code: {raw_time!r}")
            meeting["time_codes"] = None
        else:
            meeting["time_codes"] = codes
    return meetings


def number(
    value: Any, field: str, issues: list[str], *, integer: bool = False
) -> int | float | None:
    if not text(value):
        return None
    if isinstance(value, bool):
        issues.append(f"Invalid {field}: {value!r}")
        return None
    try:
        parsed = float(value)
    except (ValueError, TypeError):
        issues.append(f"Invalid {field}: {value!r}")
        return None
    if not math.isfinite(parsed) or parsed < 0 or (integer and not parsed.is_integer()):
        issues.append(f"Invalid {field}: {value!r}")
        return None
    return int(parsed) if parsed.is_integer() else parsed


def structured_fields(
    values: dict[str, Any], origins: dict[str, dict[str, Any]] | None = None
) -> dict[str, Any]:
    issues: list[str] = []
    identifier = text(values.get("id"))
    match = re.fullmatch(r"(\d{5})([A-Z]+)\s*(\d{6})", identifier)
    department = text(values.get("department_code"))
    if not department and match:
        department = match[2]
    if match and department and department != match[2]:
        issues.append("Source department code disagrees with course ID")
    lecturers = []
    for line in raw_text(values.get("lecturer")).splitlines():
        if not line.strip():
            continue
        if "\t" in line:
            names = line.split("\t")
            if len(names) > 2:
                issues.append("Unexpected extra lecturer columns")
            lecturers.append(
                {
                    "name_zh": display_text(names[0]) or None,
                    "name_en": display_text(names[1]) or None,
                }
            )
        else:
            # Historical spreadsheets use commas between Chinese names, but an
            # English surname's comma is not a person separator.
            names = (
                re.split(r"[,，、]\s*", line)
                if not re.search("[A-Za-z]", line)
                else [line]
            )
            lecturers.extend(
                {"label": display_text(name)} for name in names if name.strip()
            )
    if text(values.get("lecturer_names_en")):
        chinese_names = [
            display_text(name)
            for name in text(values.get("lecturer_names_zh")).split(",")
        ]
        english_names = [
            display_text(name) for name in text(values["lecturer_names_en"]).split(";")
        ]
        existing_names = [
            display_text(name) for name in text(values.get("lecturer")).split(",")
        ]
        if chinese_names == existing_names and len(chinese_names) == len(english_names):
            lecturers = [
                {"name_zh": zh or None, "name_en": en or None}
                for zh, en in zip(chinese_names, english_names, strict=True)
            ]
        else:
            issues.append("Teacher supplement cannot be paired with primary names")
    expertises = []
    for item in re.split(r"[\t\r\n]+", text(values.get("expertise"))):
        if not item.strip():
            continue
        expertise = re.fullmatch(r"(.*?)\s*[（(](第[一二]專長)[)）]", item.strip())
        expertises.append(
            {"name": expertise[1].strip(), "type": expertise[2]}
            if expertise
            else {"name": display_text(item), "type": None}
        )
    requirements = []
    for item in re.split(r"[\t\r\n,，]+", text(values.get("required_optional_note"))):
        if not item.strip():
            continue
        requirement = re.fullmatch(r"(.+?)\s*(必修|選修)", item.strip())
        if requirement is None:
            issues.append(f"Unrecognized curriculum requirement: {item.strip()!r}")
        requirements.append(
            {"target": requirement[1].strip(), "requirement": requirement[2]}
            if requirement
            else {"target": display_text(item), "requirement": None}
        )
    capacity = values.get("size_limit")
    cancelled = values.get("suspend")
    if cancelled is not None and text(cancelled) not in ("", "停開"):
        issues.append(f"Unrecognized cancellation marker: {cancelled!r}")
    if text(values.get("language")) not in ("", "中", "英"):
        issues.append(f"Unrecognized teaching language: {values['language']!r}")
    capacity_number = number(capacity, "size_limit", issues, integer=True)
    result = {
        "semester_code": identifier[:5],
        "department_code": department or None,
        "course_number": match[3] if match else None,
        "department_name": text(values.get("department_name")) or None,
        "college_name": text(values.get("college_name")) or None,
        "course_level_code": text(values.get("course_level_code")) or None,
        "course_level": text(values.get("course_level")) or None,
        "remote_teaching": text(values.get("remote_teaching")) or None,
        "credits": number(values.get("credit"), "credit", issues),
        "capacity": capacity_number,
        "capacity_unlimited": (
            not bool(text(capacity))
            if "size_limit" in values
            and (not text(capacity) or capacity_number is not None)
            else None
        ),
        "enrolled_count": number(
            values.get("student_count"), "student_count", issues, integer=True
        ),
        "enrollment_after_selection": number(
            values.get("enrollment_after_selection"),
            "enrollment_after_selection",
            issues,
            integer=True,
        ),
        "enrollment_after_withdrawal": number(
            values.get("enrollment_after_withdrawal"),
            "enrollment_after_withdrawal",
            issues,
            integer=True,
        ),
        "reserved_for_freshmen": number(
            values.get("freshman_reservation"),
            "freshman_reservation",
            issues,
            integer=True,
        ),
        "is_cancelled": text(cancelled) == "停開"
        if cancelled is not None and text(cancelled) in ("", "停開")
        else None,
        "language_code": {"中": "zh", "英": "en"}.get(text(values.get("language"))),
        "lecturers": lecturers,
        "meetings": parse_meetings(values, issues),
        "expertises": expertises,
        "programs": [
            display_text(item)
            for item in text(values.get("program")).split("/")
            if item.strip()
        ],
        "curriculum_requirements": requirements,
        "prerequisite_text": re.sub(
            r"\s+",
            " ",
            re.sub(r"<[^>]*>", " ", html.unescape(text(values.get("prerequisite")))),
        ).strip(),
        "parsing_issues": issues,
    }
    if match is None:
        issues.append("Course ID does not match the documented full course format")
    if (
        text(values.get("source_semester"))
        and text(values["source_semester"]) != result["semester_code"]
    ):
        issues.append("Source semester disagrees with course ID")
    if issues:
        source_rows = (
            list({json.dumps(origin, sort_keys=True) for origin in origins.values()})
            if origins
            else None
        )
        logger.warning(
            "Course %s: %s (sources=%s)",
            identifier,
            "; ".join(issues),
            sorted(source_rows) if source_rows else None,
        )
    return result


def normalize_course(
    values: dict[str, Any], origins: dict[str, dict[str, Any]] | None = None
) -> dict[str, Any]:
    identifier = text(values.get("id"))
    course_key(identifier)
    if not (text(values.get("chinese_title")) or text(values.get("english_title"))):
        raise ValueError(
            f"Course {identifier!r} has no usable course identity or title"
        )
    result = {field: display_text(values.get(field)) for field in PUBLIC_TEXT_FIELDS}
    result["id"] = identifier
    result.update(structured_fields(values, origins))
    validate_standard_course(result)
    return result


def group_courses(data: Any) -> dict[str, list[dict[str, Any]]]:
    if not isinstance(data, list) or not data:
        raise ValueError("Course collection must be a non-empty list")
    semesters: dict[str, list[dict[str, Any]]] = {}
    for index, row in enumerate(data):
        if not isinstance(row, dict):
            raise ValueError(f"Course record {index} must be an object")
        course = normalize_course(canonical_fields(row))
        semesters.setdefault(course["semester_code"], []).append(course)
    return semesters


def write_source_atomic(body: bytes, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=path.parent, prefix=f".{path.name}.", suffix=".tmp", delete=False
        ) as file:
            temporary = Path(file.name)
            file.write(body)
            file.flush()
            os.fsync(file.fileno())
        os.replace(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def load_archives(
    folder: Path,
    *,
    require_all: bool = False,
    include_latest: bool = True,
    replaced_sources: set[str] | None = None,
) -> dict[str, bytes]:
    from nthu_scraper.storage import read_json_optional

    manifest = read_json_optional(folder / "sources.json")
    previous = (
        {entry["id"]: entry for entry in manifest["sources"]}
        if manifest is not None
        else {}
    )
    inputs = {}
    replaced_sources = replaced_sources or set()
    archived_ids = {
        entry["id"]
        for entry in previous.values()
        if ARCHIVED_LATEST_ID.fullmatch(entry["id"])
    }
    archived_ids.update(
        path.stem
        for path in (folder / "raw").glob("latest-*.json")
        if ARCHIVED_LATEST_ID.fullmatch(path.stem)
    )
    for source in (*SOURCES, *(source_for_id(key) for key in sorted(archived_ids))):
        if (source.id == "latest" and not include_latest) or (
            source.id in replaced_sources
        ):
            continue
        path = folder / "raw" / source.filename
        if not path.exists():
            if require_all or source.id in previous:
                raise FileNotFoundError(
                    f"Missing course archive {path}; restore it with course_tool import"
                )
            continue
        body = path.read_bytes()
        if (
            source.id in previous
            and hashlib.sha256(body).hexdigest() != previous[source.id]["sha256"]
        ):
            raise ValueError(
                f"Course archive checksum mismatch: {path}; restore it with course_tool import"
            )
        inputs[source.id] = body
    return inputs


def load_retained_snapshots(folder: Path) -> dict[str, bytes]:
    from nthu_scraper.storage import read_json_optional

    manifest = read_json_optional(folder / "sources.json")
    previous = {
        entry["semester"]: entry
        for entry in (manifest or {}).get("retained_published_snapshots", [])
    }
    identities = set(previous)
    identities.update(
        path.stem.removeprefix("retained-")
        for path in (folder / "raw").glob("retained-*.json")
    )
    snapshots = {}
    for semester in sorted(identities):
        if not re.fullmatch(r"\d{5}", semester):
            raise ValueError(f"Invalid retained semester identity: {semester}")
        path = folder / "raw" / f"retained-{semester}.json"
        if not path.exists():
            raise FileNotFoundError(f"Missing retained published snapshot: {path}")
        body = path.read_bytes()
        if semester in previous and (
            hashlib.sha256(body).hexdigest() != previous[semester]["sha256"]
        ):
            raise ValueError(f"Retained published snapshot checksum mismatch: {path}")
        snapshots[semester] = body
    return snapshots


def validate_retained_snapshot(body: bytes, semester: str) -> int:
    try:
        rows = json.loads(body.decode("utf-8"))
        valid = (
            isinstance(rows, list)
            and bool(rows)
            and all(
                isinstance(row, dict)
                and course_key(row.get("id"))[:5] == semester
                and bool(
                    text(row.get("chinese_title")) or text(row.get("english_title"))
                )
                for row in rows
            )
        )
    except (UnicodeError, ValueError, TypeError) as error:
        raise ValueError(f"Malformed retained published snapshot {semester}") from error
    if not valid:
        raise ValueError(f"Malformed retained published snapshot {semester}")
    return len(rows)


def normalize_retained_snapshot(body: bytes, semester: str) -> list[dict[str, Any]]:
    validate_retained_snapshot(body, semester)
    result = []
    for row in json.loads(body):
        values = {
            key: row[key] for key in (*FIELD_MAPPING, *EXTRA_FIELDS) if key in row
        }
        course = normalize_course(values)
        # Published snapshots may already contain structured values that cannot
        # be recovered from their legacy display strings.
        for field in course.keys() - set(PUBLIC_TEXT_FIELDS):
            if field in row:
                if field == "parsing_issues":
                    if not isinstance(row[field], list) or not all(
                        isinstance(issue, str) for issue in row[field]
                    ):
                        raise ValueError(
                            f"Invalid retained parsing issues: {course['id']}"
                        )
                    course[field] = list(dict.fromkeys([*row[field], *course[field]]))
                else:
                    course[field] = row[field]
        validate_standard_course(course)
        result.append(course)
    return result


def validate_standard_course(course: dict[str, Any]) -> None:
    nullable_text = (
        "department_code",
        "course_number",
        "department_name",
        "college_name",
        "course_level_code",
        "course_level",
        "remote_teaching",
        "language_code",
    )
    numeric = (
        "credits",
        "capacity",
        "enrolled_count",
        "enrollment_after_selection",
        "enrollment_after_withdrawal",
        "reserved_for_freshmen",
    )
    boolean = ("capacity_unlimited", "is_cancelled")
    collections = {
        "lecturers",
        "meetings",
        "expertises",
        "programs",
        "curriculum_requirements",
        "parsing_issues",
    }
    expected = {
        *PUBLIC_TEXT_FIELDS,
        *nullable_text,
        *numeric,
        *boolean,
        *collections,
        "semester_code",
        "prerequisite_text",
    }
    valid = set(course) == expected
    valid = valid and all(
        isinstance(course[field], str)
        for field in (*PUBLIC_TEXT_FIELDS, "semester_code", "prerequisite_text")
    )
    valid = valid and all(
        course[field] is None or isinstance(course[field], str)
        for field in nullable_text
    )
    valid = valid and all(
        course[field] is None
        or (
            type(course[field]) in (int, float)
            and math.isfinite(course[field])
            and course[field] >= 0
            and (field == "credits" or float(course[field]).is_integer())
        )
        for field in numeric
    )
    valid = valid and all(
        course[field] is None or type(course[field]) is bool for field in boolean
    )
    valid = valid and all(isinstance(course[field], list) for field in collections)
    if not valid:
        raise ValueError(f"Invalid standard course: {course.get('id')!r}")
    for field in ("programs", "parsing_issues"):
        if not all(isinstance(item, str) for item in course[field]):
            raise ValueError(f"Invalid standard course {field}: {course['id']}")
    shapes = {
        "lecturers": ({"label"}, {"name_zh", "name_en"}),
        "meetings": ({"classroom", "time", "time_codes"},),
        "expertises": ({"name", "type"},),
        "curriculum_requirements": ({"target", "requirement"},),
    }
    for field, allowed in shapes.items():
        for item in course[field]:
            if not isinstance(item, dict) or set(item) not in allowed:
                raise ValueError(f"Invalid standard course {field}: {course['id']}")
            for key, value in item.items():
                if key == "time_codes":
                    valid = value is None or (
                        isinstance(value, list)
                        and all(isinstance(code, str) for code in value)
                    )
                else:
                    valid = value is None or isinstance(value, str)
                if not valid:
                    raise ValueError(f"Invalid standard course {field}: {course['id']}")
    if course["semester_code"] != course_key(course["id"])[:5]:
        raise ValueError(f"Invalid standard course semester: {course['id']}")


def prepare_converted_sources(
    inputs: dict[str, bytes], folder: Path, previous: dict[str, Any] | None
) -> dict[str, dict[str, Any]]:
    from nthu_scraper.storage import read_json_optional

    previous_sources = {
        entry["id"]: entry for entry in (previous or {}).get("sources", [])
    }
    documents = {}
    for source_id, body in sorted(inputs.items()):
        source = source_for_id(source_id)
        path = folder / "raw" / "converted" / f"{source_id}.json"
        entry = previous_sources.get(source_id, {})
        converted_entry = entry.get("converted")
        cached = read_json_optional(path)
        if cached is None and path.exists():
            raise ValueError(f"{source_id}: invalid converted course document: {path}")
        if cached is None and converted_entry is not None:
            raise FileNotFoundError(f"Missing converted course archive: {path}")
        if cached is not None:
            if converted_entry is not None and (
                hashlib.sha256(path.read_bytes()).hexdigest()
                != converted_entry["sha256"]
            ):
                raise ValueError(f"Converted course archive checksum mismatch: {path}")
            converted_records(cached, source)
            original_path = folder / "raw" / source.filename
            previous_checksum = entry.get("sha256")
            if previous_checksum is None and original_path.exists():
                previous_checksum = hashlib.sha256(
                    original_path.read_bytes()
                ).hexdigest()
            if cached["sha256"] != (
                previous_checksum or hashlib.sha256(body).hexdigest()
            ):
                raise ValueError(f"Converted course source checksum mismatch: {path}")
        documents[source_id] = (
            cached
            if cached is not None
            and cached["sha256"] == hashlib.sha256(body).hexdigest()
            else convert_source(body, source)
        )
    return documents


def publish_sources(
    inputs: dict[str, bytes],
    folder: Path,
) -> dict[str, Any]:
    from nthu_scraper.storage import read_json_optional, write_json_atomic

    inputs = dict(inputs)
    previous_manifest = read_json_optional(folder / "sources.json")
    for source_id, body in load_archives(
        folder,
        replaced_sources=set(inputs),
    ).items():
        inputs.setdefault(source_id, body)
    outgoing = folder / "raw" / "latest.json"
    if "latest" in inputs and outgoing.exists():
        previous_body = outgoing.read_bytes()
        previous_latest = (
            next(
                (
                    entry
                    for entry in previous_manifest["sources"]
                    if entry["id"] == "latest"
                ),
                None,
            )
            if previous_manifest is not None
            else None
        )
        if previous_latest is not None and (
            hashlib.sha256(previous_body).hexdigest() != previous_latest["sha256"]
        ):
            raise ValueError(
                f"Course archive checksum mismatch: {outgoing}; restore the original archive"
            )
        if previous_body != inputs["latest"]:
            outgoing_semesters = {
                course_key(row.get("科號"))[:5] for row in decode_json(previous_body)
            }
            current_semesters = {
                course_key(row.get("科號"))[:5] for row in decode_json(inputs["latest"])
            }
            if outgoing_semesters - current_semesters:
                archive_id = "latest-" + "-".join(
                    sorted(outgoing_semesters - current_semesters)
                )
                existing = inputs.get(archive_id)
                if existing is not None and existing != previous_body:
                    raise ValueError(
                        f"Conflicting course archive {archive_id}; keep the existing "
                        "archive and investigate the previous latest response"
                    )
                inputs[archive_id] = previous_body
    documents = prepare_converted_sources(inputs, folder, previous_manifest)
    records, manifest = merge_converted_sources(documents)
    semesters: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        course = normalize_course(record["values"], record["field_sources"])
        semesters.setdefault(course["semester_code"], []).append(course)
    if not semesters:
        raise ValueError("No courses to publish")
    manifest_path = folder / "sources.json"
    manifest["sources"].sort(key=lambda entry: entry["id"])
    retained = load_retained_snapshots(folder)
    for path in sorted((folder / "semesters").glob("*.json")):
        semester = path.stem
        if not re.fullmatch(r"\d{5}", semester):
            raise ValueError(f"Unexpected course semester file: {path}")
        if semester in semesters:
            continue
        body = path.read_bytes()
        if semester in retained and (
            json.loads(body)
            != normalize_retained_snapshot(retained[semester], semester)
            and body != retained[semester]
        ):
            raise ValueError(
                f"Published semester differs from retained snapshot: {path}"
            )
        retained.setdefault(semester, body)
    retained_entries = []
    for semester, body in sorted(retained.items()):
        retained_entries.append(
            {
                "id": f"retained-{semester}",
                "semester": semester,
                "path": f"raw/retained-{semester}.json",
                "sha256": hashlib.sha256(body).hexdigest(),
                "records": validate_retained_snapshot(body, semester),
                **({"superseded": True} if semester in semesters else {}),
            }
        )
    manifest["retained_published_snapshots"] = retained_entries
    manifest["semesters"] = sorted(
        {semester for entry in manifest["sources"] for semester in entry["semesters"]}
        | set(retained)
    )
    # Coverage is inferred from actual course IDs, not attachment labels.
    available = set(manifest["semesters"])
    years = [int(semester[:3]) for semester in available]
    manifest["missing_regular_semesters"] = [
        code
        for year in range(min(years), max(years) + 1)
        for term in (1, 2)
        if (code := f"{year:03}{term}0") < max(available) and code not in available
    ]
    manifest["normalized_records"] = sum(len(rows) for rows in semesters.values())
    manifest["retained_records"] = sum(
        entry["records"] for entry in retained_entries if not entry.get("superseded")
    )
    manifest["published_records"] = (
        manifest["normalized_records"] + manifest["retained_records"]
    )
    manifest["parsing_issue_records"] = sum(
        bool(row["parsing_issues"]) for rows in semesters.values() for row in rows
    )
    outputs = {
        folder / "semesters" / f"{semester}.json": rows
        for semester, rows in semesters.items()
    }
    for semester, body in retained.items():
        if semester not in semesters:
            outputs[folder / "semesters" / f"{semester}.json"] = (
                normalize_retained_snapshot(body, semester)
            )
    for source_id, document in documents.items():
        path = folder / "raw" / "converted" / f"{source_id}.json"
        outputs[path] = document
        entry = next(entry for entry in manifest["sources"] if entry["id"] == source_id)
        entry["converted"] = {
            "path": f"raw/converted/{source_id}.json",
            "sha256": hashlib.sha256(
                json.dumps(document, ensure_ascii=False, indent=4).encode("utf-8")
            ).hexdigest(),
            "schema_version": CONVERTED_VERSION,
        }
    outputs[manifest_path] = manifest
    latest = (
        [
            normalize_course(canonical_fields(record.values))
            for record in converted_records(documents["latest"], SOURCE_BY_ID["latest"])
        ]
        if "latest" in documents
        else None
    )
    if latest is not None:
        outputs[folder / "latest.json"] = latest
    # Check every existing JSON and serialize every new result before any write.
    for path, value in outputs.items():
        read_json_optional(path)
        json.dumps(value, ensure_ascii=False, allow_nan=False)
    for source_id, body in sorted(inputs.items()):
        path = folder / "raw" / source_for_id(source_id).filename
        if not path.exists() or path.read_bytes() != body:
            write_source_atomic(body, path)
    for semester, body in sorted(retained.items()):
        archived = folder / "raw" / f"retained-{semester}.json"
        if not archived.exists():
            write_source_atomic(body, archived)
    for path, value in outputs.items():
        indent = 2 if path == folder / "latest.json" else 4
        serialized = json.dumps(value, ensure_ascii=False, indent=indent)
        if path.parent == folder / "raw" / "converted":
            body = serialized.encode("utf-8")
            if not path.exists() or path.read_bytes() != body:
                write_source_atomic(body, path)
        elif not path.exists() or path.read_text(encoding="utf-8") != serialized:
            write_json_atomic(value, path, indent=indent)
    (folder.parent / "courses.json").unlink(missing_ok=True)
    logger.info(
        "Published courses: %s; missing regular semesters=%s; filled fields=%s; "
        "conflicts=%s; records with parsing issues=%s",
        {semester: len(rows) for semester, rows in semesters.items()},
        manifest["missing_regular_semesters"],
        manifest["filled_fields"],
        manifest["conflicts"],
        manifest["parsing_issue_records"],
    )
    return manifest

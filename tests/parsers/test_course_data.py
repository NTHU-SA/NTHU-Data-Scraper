import hashlib
import json
from io import BytesIO

import pytest
import scrapy
from openpyxl import Workbook
from scrapy.http import TextResponse
from twisted.internet.error import TimeoutError
from twisted.python.failure import Failure

from nthu_scraper.course_data import (
    FIELD_MAPPING,
    SOURCE_BY_ID,
    canonical_fields,
    decode_json,
    group_courses,
    merge_sources,
    parse_meetings,
    publish_sources,
    read_source,
    structured_fields,
)
from nthu_scraper.spiders import nthu_courses
from nthu_scraper.spiders.nthu_courses import CoursesSpider

ID = "11020SOC 578900"
MAPPING = FIELD_MAPPING


def encoded(rows):
    return json.dumps(rows, ensure_ascii=False).encode("utf-8")


def workbook_bytes(headers, rows, *, sheet_name="108-110", teachers=None):
    workbook = Workbook()
    workbook.active.title = "下載日期"
    workbook.active.append(["Not a course sheet"])
    sheet = workbook.create_sheet(sheet_name)
    sheet.append(["A preamble, not the header"])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if teachers is not None:
        sheet = workbook.create_sheet("1141110教師data")
        sheet.append(["科號", "課程中文名稱", "授課教師中文", "授課教師英文"])
        for row in teachers:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_raw_json_encoding_and_wrappers():
    data = [{"科號": ID, "中文課名": "測試"}]
    assert decode_json(encoded({"工作表1": data})) == data
    assert decode_json(json.dumps(data).encode("utf-16")) == data
    assert decode_json(b"\xef\xbb\xbf" + encoded(data)) == data
    for invalid in ([], {}, [None]):
        with pytest.raises(ValueError):
            decode_json(encoded(invalid))


def test_workbook_uses_configured_sheet_and_header_not_first_sheet():
    body = workbook_bytes(
        ["科號", "中文課名", "開課單位", "開課院別", "人限", "專長相關"],
        [[ID, "測試", "社會學研究所", "人文社會學院", None, "社會學(第二專長)"]],
    )
    records = read_source(body, SOURCE_BY_ID["108-110"])
    assert len(records) == 1
    assert records[0].origin == {"source_id": "108-110", "sheet": "108-110", "row": 3}
    fields = canonical_fields(records[0].values, MAPPING)
    assert fields["department_name"] == "社會學研究所"
    assert fields["college_name"] == "人文社會學院"
    assert fields["size_limit"] == ""
    assert fields["expertise"] == "社會學(第二專長)"


def test_merge_excel_precedence_missing_values_and_stable_order():
    excel = workbook_bytes(
        ["科號", "中文課名", "學分", "人限", "總人數", "備註", "課程列入系所班別"],
        [[ID, "Excel title", 0.5, None, 0, None, "社會所110M  必修"]],
    )
    raw = encoded(
        [
            {
                "科號": ID,
                "中文課名": "JSON title",
                "學分": 3,
                "人限": "50",
                "總人數": 99,
                "備註": "補回",
                "授課語言": "英",
            },
            {"科號": "11020SOC 580200", "中文課名": "Only JSON"},
        ]
    )
    inputs = {"10910-11110": raw, "108-110": excel}
    records, report = merge_sources(inputs, MAPPING)
    assert (records, report) == merge_sources(
        dict(reversed(list(inputs.items()))), MAPPING
    )
    first = records[0]
    assert first["values"]["chinese_title"] == "Excel title"
    assert first["values"]["credit"] == 0.5
    assert first["values"]["size_limit"] == ""
    assert first["values"]["student_count"] == 0
    assert first["values"]["note"] == "補回"
    assert first["field_sources"]["note"]["source_id"] == "10910-11110"
    assert first["field_sources"]["required_optional_note"]["source_id"] == "108-110"
    assert len(records) == 2
    assert report["conflicts"]["chinese_title"] == 1


def test_duplicate_conflicts_are_explicit():
    rows = [
        {"科號": ID, "中文課名": "A"},
        {"科號": ID, "中文課名": "B"},
    ]
    with pytest.raises(ValueError, match="duplicate"):
        merge_sources({"latest": encoded(rows)}, MAPPING)
    with pytest.raises(ValueError, match="equal-priority"):
        merge_sources(
            {"10820": encoded(rows[:1]), "10910-11110": encoded(rows[1:])}, MAPPING
        )


def test_split_converter_row_requires_intact_excel_recovery(caplog):
    raw = encoded(
        [
            {"科號": ID, "中文課名": "測試", "英文課名": '"Title, with commas'},
            {"科號": '"', "中文課名": "3"},
        ]
    )
    with pytest.raises(ValueError, match="unrecoverable"):
        merge_sources({"10910-11110": raw}, MAPPING)
    excel = workbook_bytes(
        ["科號", "中文課名", "英文課名", "學分"],
        [[ID, "測試", "Title, with commas\n", 3]],
    )
    records, report = merge_sources({"10910-11110": raw, "108-110": excel}, MAPPING)
    assert len(records) == 1
    assert records[0]["values"]["english_title"] == "Title, with commas\n"
    assert len(report["source_repairs"]) == 1
    assert report["source_repairs"][0]["row"] == 2
    assert "Recovered malformed" in caplog.text


def test_multiple_meetings_preserve_all_pairs_without_legacy_fields():
    raw = {
        "科號": "11510AES 520300",
        "中文課名": "環境化學",
        "教室與上課時間": "BMES醫環618\tM3M4\nBMES醫環501\tF3F4\n",
    }
    result = group_courses([raw])["11510"][0]
    assert "classroom" not in result
    assert "time" not in result
    assert result["meetings"] == [
        {"classroom": "BMES醫環618", "time": "M3M4", "time_codes": ["M3", "M4"]},
        {"classroom": "BMES醫環501", "time": "F3F4", "time_codes": ["F3", "F4"]},
    ]
    assert result["department_code"] == "AES"
    assert result["course_number"] == "520300"


def test_historical_comma_meetings_and_incomplete_pairs():
    issues = []
    result = parse_meetings({"classroom": "化 326,化 324", "time": "M5,M6M7M8"}, issues)
    assert [(row["classroom"], row["time_codes"]) for row in result] == [
        ("化 326", ["M5"]),
        ("化 324", ["M6", "M7", "M8"]),
    ]
    assert issues == []
    assert parse_meetings({"time": "SnSa"}, issues) == [
        {"classroom": None, "time": "SnSa", "time_codes": ["Sn", "Sa"]}
    ]
    parse_meetings({"classroom": "A,B", "time": "M1,M2,M3"}, issues)
    assert issues == ["Unequal comma-separated classroom/time counts"]


def test_empty_rooms_do_not_shift_times_or_consume_tab_delimiters():
    issues = []
    assert parse_meetings({"class_room_and_time": "\tM3M4\n"}, issues) == [
        {"classroom": None, "time": "M3M4", "time_codes": ["M3", "M4"]}
    ]
    assert parse_meetings({"class_room_and_time": "Room\t\n"}, issues) == [
        {"classroom": "Room", "time": None, "time_codes": []}
    ]
    assert parse_meetings({"classroom": "\nB", "time": "M1\nT2"}, issues) == [
        {"classroom": None, "time": "M1", "time_codes": ["M1"]},
        {"classroom": "B", "time": "T2", "time_codes": ["T2"]},
    ]
    assert issues == []


@pytest.mark.parametrize(
    ("override", "expected"),
    [
        ({"classroom": "New"}, [("New", "M1")]),
        ({"time": "T2"}, [("Old", "T2")]),
        ({"classroom": ""}, [(None, "M1")]),
        ({"time": ""}, [("Old", None)]),
    ],
)
def test_partial_meeting_overrides_preserve_missing_component(override, expected):
    issues = []
    meetings = parse_meetings({"class_room_and_time": "Old\tM1", **override}, issues)
    assert [(item["classroom"], item["time"]) for item in meetings] == expected
    assert issues == []


@pytest.mark.parametrize(
    ("override", "expected"),
    [
        (
            {"classroom": "New\nOther"},
            [("New", "M1"), ("Other", "T2")],
        ),
        (
            {"time": "W3\nR4"},
            [("Old", "W3"), ("Next", "R4")],
        ),
        (
            {"classroom": ""},
            [(None, "M1"), (None, "T2")],
        ),
        (
            {"time": ""},
            [("Old", None), ("Next", None)],
        ),
    ],
)
def test_multiple_partial_meeting_overrides_align_pairs(override, expected):
    raw = {"科號": ID, "中文課名": "測試", "教室與上課時間": "Old\tM1\nNext\tT2"}
    if "classroom" in override:
        raw["教室"] = override["classroom"]
    if "time" in override:
        raw["上課時間"] = override["time"]
    result = group_courses([raw])["11020"][0]
    meetings = result["meetings"]
    assert [(item["classroom"], item["time"]) for item in meetings] == expected
    assert "classroom" not in result
    assert "time" not in result
    assert result["parsing_issues"] == []


def test_partial_meeting_count_mismatch_keeps_raw_and_reports_alignment():
    raw = {
        "科號": ID,
        "中文課名": "測試",
        "教室與上課時間": "Old\tM1\nNext\tT2",
        "教室": "New",
    }
    result = group_courses([raw])["11020"][0]
    assert [(m["classroom"], m["time"]) for m in result["meetings"]] == [
        ("New", "M1"),
        (None, "T2"),
    ]
    assert "classroom" not in result
    assert "time" not in result
    assert result["parsing_issues"] == [
        "Unequal separate classroom/combined meeting counts"
    ]
    assert raw["教室與上課時間"] == "Old\tM1\nNext\tT2"
    assert "class_room_and_time" not in result
    time_issues = []
    assert parse_meetings(
        {"class_room_and_time": "Old\tM1", "time": "T2\nW3"}, time_issues
    )[1] == {"classroom": None, "time": "W3", "time_codes": ["W3"]}
    assert time_issues == ["Unequal separate time/combined meeting counts"]


def test_partial_override_aligns_comma_separated_rooms_to_combined_times():
    issues = []
    meetings = parse_meetings(
        {"class_room_and_time": "Old\tM1\nNext\tT2", "classroom": "A,B"},
        issues,
    )
    assert [(m["classroom"], m["time"]) for m in meetings] == [
        ("A", "M1"),
        ("B", "T2"),
    ]
    assert issues == []


def test_lists_and_numeric_semantics(caplog):
    result = structured_fields(
        {
            "id": "11510ANTH500000",
            "lecturer": "陳俊銘\tCHEN, CHUN-MING\n周秀專\tCHOU, HSIU-CHUAN\n",
            "expertise": "環境科技學程(第二專長)\t能源與環境(第一專長)",
            "program": "學程 A/學程 B",
            "required_optional_note": "分環所115M  必修\t原科院學士班112B  選修",
            "credit": "0.3",
            "size_limit": "",
            "freshman_reservation": "0",
            "suspend": "",
            "prerequisite": "<div>A &amp; B<br/> C</div>",
        }
    )
    assert result["department_code"] == "ANTH"
    assert result["lecturers"][0] == {"name_zh": "陳俊銘", "name_en": "CHEN, CHUN-MING"}
    assert result["expertises"][1] == {"name": "能源與環境", "type": "第一專長"}
    assert result["programs"] == ["學程 A", "學程 B"]
    assert result["curriculum_requirements"][0] == {
        "target": "分環所115M",
        "requirement": "必修",
    }
    assert result["credits"] == 0.3
    assert result["capacity"] is None
    assert result["capacity_unlimited"] is True
    assert result["reserved_for_freshmen"] == 0
    assert result["is_cancelled"] is False
    assert result["enrolled_count"] is None
    assert result["prerequisite_text"] == "A & B C"
    assert result["parsing_issues"] == []
    assert structured_fields({"id": "11510ANTH500000"})["is_cancelled"] is None
    zero = structured_fields({"id": "11510ANTH500000", "size_limit": "0"})
    assert zero["capacity"] == 0
    assert zero["capacity_unlimited"] is False
    invalid = structured_fields(
        {"id": "11510ANTH500000", "student_count": "#N/A", "suspend": "unknown"}
    )
    assert invalid["enrolled_count"] is None
    assert invalid["is_cancelled"] is None
    assert len(invalid["parsing_issues"]) == 2
    assert "Invalid student_count" in caplog.text
    compact = structured_fields(
        {
            "id": "11020PHYS101001",
            "required_optional_note": "物理系109B B必修,物理系109B A選修",
        }
    )
    assert compact["curriculum_requirements"] == [
        {"target": "物理系109B B", "requirement": "必修"},
        {"target": "物理系109B A", "requirement": "選修"},
    ]
    assert compact["parsing_issues"] == []
    invalid_capacity = structured_fields(
        {"id": "11020PHYS101001", "size_limit": "not a number"}
    )
    assert invalid_capacity["capacity_unlimited"] is None


def test_teacher_supplement_joins_only_main_courses_and_tracks_origin():
    body = workbook_bytes(
        ["科號", "中文課名", "教師"],
        [["11410AES 470200", "測試", "陳俊銘, 周秀專"]],
        sheet_name="11410",
        teachers=[
            [
                "11410AES 470200",
                "測試",
                "陳俊銘, 周秀專",
                "CHEN, CHUN-MING; CHOU, HSIU-CHUAN",
            ],
            ["11410AES 999999", "Not in primary sheet", "其他", "OTHER"],
        ],
    )
    records, _ = merge_sources({"11410": body}, MAPPING)
    assert len(records) == 1
    result = structured_fields(records[0]["values"])
    assert result["lecturers"][1] == {
        "name_zh": "周秀專",
        "name_en": "CHOU, HSIU-CHUAN",
    }
    assert (
        records[0]["field_sources"]["lecturer_names_en"]["sheet"] == "1141110教師data"
    )
    assert len(records[0]["sources"]) == 2


def test_publication_preserves_raw_and_is_byte_stable(tmp_path):
    raw = encoded(
        [
            {"科號": "11410AES 450100", "中文課名": "歷史"},
            {"科號": "11510AES 450100", "中文課名": "最新"},
        ]
    )
    folder = tmp_path / "courses"
    root = tmp_path / "courses.json"
    report = publish_sources({"latest": raw}, folder)
    assert not root.exists()
    assert report["missing_regular_semesters"] == ["11420"]
    assert (folder / "raw" / "latest.json").read_bytes() == raw
    assert report["sources"][0]["sha256"] == hashlib.sha256(raw).hexdigest()
    before = {
        p: (p.read_bytes(), p.stat().st_mtime_ns)
        for p in tmp_path.rglob("*")
        if p.is_file()
    }
    publish_sources({"latest": raw}, folder)
    assert before == {p: (p.read_bytes(), p.stat().st_mtime_ns) for p in before}
    with pytest.raises(ValueError):
        publish_sources({"latest": encoded([{}])}, folder)
    assert before == {p: (p.read_bytes(), p.stat().st_mtime_ns) for p in before}


def test_corrupt_any_target_prevents_all_source_and_output_writes(tmp_path):
    folder = tmp_path / "courses"
    folder.mkdir()
    root = folder / "latest.json"
    root.write_text("broken", encoding="utf-8")
    with pytest.raises(json.JSONDecodeError):
        publish_sources(
            {"latest": encoded([{"科號": "11510AES 450100", "中文課名": "測試"}])},
            folder,
        )
    assert not (folder / "raw").exists()
    assert not (folder / "semesters").exists()


def test_spider_downloads_only_latest():
    start = CoursesSpider().start()
    with pytest.raises(StopIteration) as first:
        start.__anext__().send(None)
    assert first.value.value.url == SOURCE_BY_ID["latest"].url
    with pytest.raises(StopAsyncIteration):
        start.__anext__().send(None)


@pytest.mark.parametrize("mode", ["backfill", "rebuild", "invalid"])
def test_offline_operations_are_not_spider_modes(mode):
    with pytest.raises(ValueError, match="course_tool"):
        CoursesSpider(mode=mode)


def test_latest_refresh_reuses_archives_without_losing_history(tmp_path, monkeypatch):
    folder = tmp_path / "courses"
    root = folder / "latest.json"
    monkeypatch.setattr(nthu_courses, "COURSES_FOLDER", folder)
    body = workbook_bytes(
        ["科號", "中文課名", "開課單位"], [[ID, "歷史", "社會學研究所"]]
    )
    raw = folder / "raw"
    raw.mkdir(parents=True)
    (raw / "108-110.xlsx").write_bytes(body)
    latest = encoded([{"科號": "11510AES 450100", "中文課名": "最新"}])
    response = TextResponse(
        SOURCE_BY_ID["latest"].url,
        body=latest,
        request=scrapy.Request(
            SOURCE_BY_ID["latest"].url, meta={"data_type": "latest"}
        ),
    )
    CoursesSpider().parse(response)
    old = json.loads((folder / "semesters" / "11020.json").read_text(encoding="utf-8"))
    assert old[0]["department_name"] == "社會學研究所"
    manifest = json.loads((folder / "sources.json").read_text(encoding="utf-8"))
    assert manifest["normalized_records"] == 2
    assert {source["id"] for source in manifest["sources"]} == {"latest", "108-110"}
    before = root.read_bytes()
    (raw / "108-110.xlsx").write_bytes(b"broken archive")
    with pytest.raises(ValueError, match="checksum mismatch"):
        CoursesSpider().parse(response)
    assert root.read_bytes() == before
    (raw / "108-110.xlsx").unlink()
    with pytest.raises(FileNotFoundError, match="Missing course archive"):
        CoursesSpider().parse(response)
    assert root.read_bytes() == before


def test_download_failure_is_explicit_and_blocks_buffered_publication():
    spider = CoursesSpider()
    failure = Failure(TimeoutError())
    failure.request = scrapy.Request(SOURCE_BY_ID["latest"].url)
    with pytest.raises(ValueError, match="Course download failed"):
        spider.handle_request_error(failure)
    assert spider.crawl_incomplete


def test_two_stage_excel_preserves_raw_columns_and_is_independently_replayable():
    from nthu_scraper.course_data import convert_source, merge_converted_sources

    body = workbook_bytes(
        ["科號", "中文課名", "教師", "備註", "Unmodeled"],
        [[ID, " Title<BR>raw ", "甲,乙", " raw\tvalue\n", "keep this"]],
    )
    document = convert_source(body, SOURCE_BY_ID["108-110"])
    document = json.loads(json.dumps(document, ensure_ascii=False))
    assert document["schema_version"] == 1
    assert document["sha256"] == hashlib.sha256(body).hexdigest()
    record = document["records"][0]
    assert record["values"]["中文課名"] == " Title<BR>raw "
    assert record["values"]["備註"] == " raw\tvalue\n"
    assert record["values"]["Unmodeled"] == "keep this"
    assert record["origin"]["row"] == 3
    records, _ = merge_converted_sources({"108-110": document})
    from nthu_scraper.course_data import normalize_course

    result = normalize_course(records[0]["values"])
    assert result["chinese_title"] == "Title raw"
    assert result["note"] == "raw value"
    assert result["lecturers"] == [{"label": "甲"}, {"label": "乙"}]
    assert "Unmodeled" not in result


def test_teacher_supplement_cannot_pair_is_archived_not_guessed(tmp_path):
    body = workbook_bytes(
        ["科號", "中文課名", "教師"],
        [["11410AES 470200", "Title", "甲,乙"]],
        sheet_name="11410",
        teachers=[["11410AES 470200", "Title", "甲,乙", "Only one name"]],
    )
    folder = tmp_path / "courses"
    publish_sources({"11410": body}, folder)
    result = json.loads(
        (folder / "semesters" / "11410.json").read_text(encoding="utf-8")
    )[0]
    assert result["lecturers"] == [{"label": "甲"}, {"label": "乙"}]
    assert result["parsing_issues"] == [
        "Teacher supplement cannot be paired with primary names"
    ]
    assert "lecturer_names_en" not in result
    document = json.loads(
        (folder / "raw" / "converted" / "11410.json").read_text(encoding="utf-8")
    )
    assert document["records"][0]["values"]["授課教師英文"] == "Only one name"
    assert (folder / "raw" / "11410.xlsx").read_bytes() == body


@pytest.mark.parametrize("damage", ["missing", "malformed", "changed"])
def test_bad_converted_archive_prevents_refresh_writes(tmp_path, damage):
    folder = tmp_path / "courses"
    old = encoded([{"科號": "11510AES 450100", "中文課名": "Old"}])
    publish_sources({"latest": old}, folder)
    path = folder / "raw" / "converted" / "latest.json"
    if damage == "missing":
        path.unlink()
    elif damage == "malformed":
        path.write_bytes(b"{")
    else:
        document = json.loads(path.read_bytes())
        document["records"][0]["values"]["中文課名"] = "Tampered"
        path.write_text(json.dumps(document), encoding="utf-8")
    legacy_root = tmp_path / "courses.json"
    legacy_root.write_bytes(b"[]")
    before = {path: path.read_bytes() for path in tmp_path.rglob("*") if path.is_file()}
    with pytest.raises((ValueError, FileNotFoundError)):
        publish_sources(
            {"latest": encoded([{"科號": "11510AES 450100", "中文課名": "New"}])},
            folder,
        )
    assert before == {path: path.read_bytes() for path in before}
